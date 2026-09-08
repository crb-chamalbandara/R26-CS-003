"""
Build the submission-ready Word (.docx) twin of ICAC2026_BeaconDetector.tex.

Same content, same section order, same tables and figures, same reference
list -- no facts are re-derived here, only re-typeset. Two-column IEEE-style
body via direct oxml section manipulation (python-docx has no first-class API
for column count), Times New Roman, numbered IEEE-style citations.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PAPER_DIR = Path(__file__).resolve().parent
FIG = PAPER_DIR / "figures"
OUT = PAPER_DIR / "ICAC2026_BeaconDetector.docx"

FONT = "Times New Roman"


def set_two_columns(section, n=2, space_twips=432):
    sectPr = section._sectPr
    cols = sectPr.find(qn('w:cols'))
    if cols is None:
        cols = OxmlElement('w:cols')
        sectPr.append(cols)
    cols.set(qn('w:num'), str(n))
    cols.set(qn('w:space'), str(space_twips))
    cols.set(qn('w:equalWidth'), "1")


def base_style(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)
    rpr = normal.element.get_or_add_rPr()
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rpr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), FONT)
    pf = normal.paragraph_format
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0


def add_heading(doc, text, level=1, numbering=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run((numbering + "  " if numbering else "") + text)
    run.font.name = FONT
    run.font.size = Pt(10 if level == 1 else 10)
    run.bold = True
    run.font.small_caps = (level == 1)
    if level == 2:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run.italic = True
        run.font.small_caps = False
    return p


def add_body(doc, text, justify=True, space_after=6, size=10, indent_first=True):
    p = doc.add_paragraph()
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(space_after)
    if indent_first:
        p.paragraph_format.first_line_indent = Inches(0.2)
    run = p.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(size)
    return p


def add_bullets(doc, items, size=10):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.22)
        run = p.add_run(it)
        run.font.name = FONT
        run.font.size = Pt(size)


def add_figure(doc, path, caption, width_in=3.4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_in))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = cap.add_run(caption)
    cr.font.name = FONT
    cr.font.size = Pt(8.5)
    cap.paragraph_format.space_after = Pt(8)


def add_table(doc, caption, headers, rows, col_widths=None, bold_cells=None):
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = cap.add_run(caption)
    cr.font.name = FONT; cr.font.size = Pt(9); cr.bold = False
    cr.font.small_caps = True

    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    bold_cells = bold_cells or set()
    for j, h in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.text = ""
        r = cell.paragraphs[0].add_run(h)
        r.font.name = FONT; r.font.size = Pt(8.5); r.bold = True
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.rows[i + 1].cells[j]
            cell.text = ""
            r = cell.paragraphs[0].add_run(str(val))
            r.font.name = FONT; r.font.size = Pt(8.5)
            r.bold = (i, j) in bold_cells
            cell.paragraphs[0].alignment = (WD_ALIGN_PARAGRAPH.LEFT if j == 0
                                            else WD_ALIGN_PARAGRAPH.CENTER)
    if col_widths:
        for row in t.rows:
            for j, w in enumerate(col_widths):
                row.cells[j].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return t


# ─────────────────────────────────────────────────────────────────────────
# Document assembly. All prose, tables and references live in the content
# module so this file stays pure formatting.
# ─────────────────────────────────────────────────────────────────────────
import icac2026_beacondetector_content as C


def title_block(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(C.TITLE)
    r.font.name = FONT; r.font.size = Pt(20); r.bold = False
    p.paragraph_format.space_after = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(C.AUTHORS)
    r.font.name = FONT; r.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(2)

    for line in C.AFFIL.split("\n"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        r.font.name = FONT; r.font.size = Pt(10)
        p.paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def abstract_block(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run("Abstract—")
    r.font.name = FONT; r.font.size = Pt(9); r.bold = True; r.italic = True
    r2 = p.add_run(C.ABSTRACT)
    r2.font.name = FONT; r2.font.size = Pt(9); r2.bold = True
    p.paragraph_format.space_after = Pt(6)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run("Index Terms—")
    r.font.name = FONT; r.font.size = Pt(9); r.bold = True; r.italic = True
    r2 = p.add_run(C.KEYWORDS)
    r2.font.name = FONT; r2.font.size = Pt(9); r2.bold = True
    p.paragraph_format.space_after = Pt(8)


def main():
    doc = Document()
    base_style(doc)
    for s in doc.sections:
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)      # A4, per ICAC
        s.top_margin = Inches(0.75); s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(0.62); s.right_margin = Inches(0.62)

    title_block(doc)

    body = doc.add_section(WD_SECTION.CONTINUOUS)
    set_two_columns(body)
    body.page_width, body.page_height = Cm(21.0), Cm(29.7)
    body.top_margin = Inches(0.75); body.bottom_margin = Inches(1.0)
    body.left_margin = Inches(0.62); body.right_margin = Inches(0.62)

    abstract_block(doc)

    add_heading(doc, "Introduction", 1, "I.")
    for para in C.INTRO:
        add_body(doc, para)
    add_bullets(doc, C.CONTRIBUTIONS)

    add_heading(doc, "Literature Review", 1, "II.")
    for letter, (subhead, para) in zip("ABCDE", C.LITREVIEW):
        add_heading(doc, subhead, 2, f"{letter}.")
        add_body(doc, para)

    add_heading(doc, "Methodology", 1, "III.")
    add_figure(doc, FIG / "fig1_architecture.png", C.FIG1_CAPTION, width_in=3.35)
    add_heading(doc, "Non-Invasive Collection", 2, "A.")
    add_body(doc, C.COLLECTION)
    add_heading(doc, "Feature Vector", 2, "B.")
    add_body(doc, C.FEATURE_VECTOR)
    add_table(doc, C.TABLE1["caption"], C.TABLE1["headers"], C.TABLE1["rows"],
              C.TABLE1["col_widths"], C.TABLE1["bold_cells"])
    add_heading(doc, "Heuristic Layer and Fusion", 2, "C.")
    for para in C.HEURISTIC:
        add_body(doc, para)
    add_heading(doc, "Dataset Construction", 2, "D.")
    add_body(doc, C.DATASET_INTRO)
    add_heading(doc, "Pseudo-Replication", 2, "E.")
    add_body(doc, C.PSEUDOREP)
    add_heading(doc, "Out-of-Scope Concept", 2, "F.")
    add_body(doc, C.OUTOFSCOPE)
    add_table(doc, C.TABLE2["caption"], C.TABLE2["headers"], C.TABLE2["rows"],
              C.TABLE2["col_widths"], C.TABLE2["bold_cells"])
    add_heading(doc, "Scope Criterion and Cap", 2, "G.")
    for para in C.SCOPE:
        add_body(doc, para)
    add_heading(doc, "Classifier and Training", 2, "H.")
    add_body(doc, C.MODEL[0])
    add_heading(doc, "Calibration and Threshold Placement", 2, "I.")
    add_body(doc, C.MODEL[1])
    add_body(doc, C.MODEL[2])
    add_heading(doc, "Evaluation Protocol", 2, "J.")
    add_body(doc, C.MODEL[3])
    add_body(doc, C.MODEL[4])

    add_heading(doc, "Results and Discussion", 1, "IV.")
    add_figure(doc, FIG / "fig2_headline_results.png", C.FIG2_CAPTION, width_in=3.35)
    add_heading(doc, "Headline Results", 2, "A.")
    for para in C.HEADLINE:
        add_body(doc, para)
    add_table(doc, C.TABLE3["caption"], C.TABLE3["headers"], C.TABLE3["rows"],
              C.TABLE3["col_widths"], C.TABLE3["bold_cells"])
    add_heading(doc, "Why the Gradient-Boosted Classifier Was Selected: A Five-Classifier Comparison", 2, "B.")
    add_body(doc, C.WHYXGB)
    add_figure(doc, FIG / "fig3_classifier_comparison.png", C.FIG3_CAPTION, width_in=3.35)
    add_table(doc, C.TABLE4["caption"], C.TABLE4["headers"], C.TABLE4["rows"],
              C.TABLE4["col_widths"], C.TABLE4["bold_cells"])
    add_heading(doc, "Fusion Weight and a Structural Limit", 2, "C.")
    for para in C.FUSION:
        add_body(doc, para)
    add_heading(doc, "Discussion", 2, "D.")
    add_body(doc, C.DISCUSSION)

    add_heading(doc, "Limitations and Future Work", 1, "V.")
    add_bullets(doc, C.LIMITATIONS)
    add_body(doc, C.FUTURE_WORK)

    add_heading(doc, "Conclusion", 1, "VI.")
    add_body(doc, C.CONCLUSION)

    add_heading(doc, "Acknowledgment", 1, "")
    add_body(doc, C.ACK)

    add_heading(doc, "References", 1, "")
    for i, ref in enumerate(C.REFERENCES, 1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.first_line_indent = Inches(-0.22)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"[{i}] {ref}")
        r.font.name = FONT; r.font.size = Pt(8)

    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
